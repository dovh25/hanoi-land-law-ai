"""Four explicitly synthetic formats, generated locally with no network or ML weights."""
from pathlib import Path
import json
from batch.contracts import canonical

TEXT = """SYNTHETIC ONLY — không phải văn bản pháp luật chính thức.
Căn cứ dữ liệu giả lập để kiểm thử.
Phần thứ nhất
Chương I
Mục 1
Tiểu mục 1
Điều 1. Quy tắc giả lập
1. Nội dung khoản giả lập.
a) Nội dung điểm giả lập.
đ) Điểm có chữ đ và dấu tiếng Việt.
2. Khoản tiếp theo không bị bỏ sót.
Điều 2. Điều giả lập tiếp theo
1. Không có giá trị pháp lý.
Phụ lục I
Bảng giả lập | Giá trị
Ghi chú và chú thích cuối tài liệu giả lập.
"""


def sources(directory):
    root = Path(directory)
    root.mkdir(parents=True, exist_ok=False)
    (root / "synthetic.txt").write_text(TEXT, encoding="utf-8")
    (root / "synthetic.html").write_text("<html><body>" + "".join(f"<p>{line}</p>" for line in TEXT.splitlines()) +
        "<table><tr><th>SYNTHETIC</th><th>Giá trị</th></tr><tr><td>A</td><td>1</td></tr></table><p>Chú thích sau bảng.</p></body></html>", encoding="utf-8")
    from docx import Document
    doc = Document()
    for line in TEXT.splitlines(): doc.add_paragraph(line)
    doc.add_table(rows=1, cols=2).rows[0].cells[0].text = "SYNTHETIC table"
    doc.save(root / "synthetic.docx")
    from pypdf import PdfWriter
    from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject
    writer = PdfWriter()
    page = writer.add_blank_page(width=612, height=792)
    font = DictionaryObject({NameObject('/Type'): NameObject('/Font'), NameObject('/Subtype'): NameObject('/Type1'), NameObject('/BaseFont'): NameObject('/Helvetica')})
    page[NameObject('/Resources')] = DictionaryObject({NameObject('/Font'): DictionaryObject({NameObject('/F1'): font})})
    stream = DecodedStreamObject()
    stream.set_data(b'BT /F1 12 Tf 50 740 Td (SYNTHETIC PDF fixture - NOT OFFICIAL LAW.) Tj 0 -20 Td (Preamble, attachment and footnote text preserved.) Tj ET')
    page[NameObject('/Contents')] = writer._add_object(stream)
    writer.write(root / "synthetic.pdf")
    manifests = []
    for suffix in ("txt", "html", "docx", "pdf"):
        metadata = dict(file=f"synthetic.{suffix}", official_identifier=f"SYNTHETIC:PHASE-C:{suffix}",
            official_number=f"SYNTHETIC-C-{suffix}", document_type="synthetic_test", title=f"SYNTHETIC {suffix} test",
            issuing_authority="SYNTHETIC authority", jurisdiction="VN-HN", source_url=f"https://synthetic.invalid/{suffix}",
            acquired_at="2026-01-01T00:00:00+00:00", is_synthetic=True)
        path = root / f"{suffix}.manifest.json"
        path.write_text(canonical(metadata) + "\n", encoding="utf-8"); manifests.append(path)
    return manifests


def smoke(workspace):
    from preprocess.offline_ingestion import initialize, ingest, review, seal, export
    root = initialize(workspace)
    if (root / "synthetic-result.json").exists():
        result = json.loads((root / "synthetic-result.json").read_text())
        from batch.files import validate_output
        validate_output(root / "embedding-input", root / "embedding-output")
        return result
    manifest_dir = root / "synthetic-sources"
    manifests = sorted(manifest_dir.glob('*.manifest.json')) if manifest_dir.exists() else sources(manifest_dir)
    stages = []
    for manifest in manifests:
        stage = ingest(root, manifest)
        if stage["status"] != "pending": raise ValueError("Synthetic extraction quarantined; inspect stage report")
        stages.append(stage["stage_id"])
        review(root, stage["stage_id"], dict(decision="approved", reviewer="SYNTHETIC operator",
            reason="Synthetic extraction reviewed, not authoritative law", evidence_locator="SYNTHETIC entire fixture",
            reviewed_at="2026-01-02T00:00:00+00:00", validity="unknown"))
    result = seal(root, "SYNTHETIC-phase-c-v1", stages, approve_publication=True, purpose="synthetic")
    if not (root / "embedding-input").exists(): export(root, result["release_hash"], root / "embedding-input")
    from batch.embedding import SyntheticEncoder
    from batch.worker import run_batch
    from batch.colab import package_output
    run_batch(root / "embedding-input", root / "embedding-output", SyntheticEncoder())
    if not (root / "embedding-output.zip").exists():
        package_output(root / "embedding-input", root / "embedding-output", root / "embedding-output.zip")
    result.update(synthetic=True, real_documents=0)
    (root / "synthetic-result.json").write_text(canonical(result) + "\n", encoding="utf-8")
    return result
